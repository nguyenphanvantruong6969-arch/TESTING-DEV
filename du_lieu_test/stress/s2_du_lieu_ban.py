# -*- coding: utf-8 -*-
"""Nhóm 2 — Dữ liệu bẩn / biên (fuzz Sổ nhập và CSV). CHỈ ĐO, không đặt ngưỡng. Dữ liệu MÔ PHỎNG.

    python du_lieu_test/stress/s2_du_lieu_ban.py [--nhanh] [--so-bien-the N]
    python du_lieu_test/stress/s2_du_lieu_ban.py --lai OP:SEED      # chạy lại đúng một biến thể

Mỗi biến thể = (toán tử, seed) → tất định. Chạy trên bản sao một CSDL nền ĐÃ có dữ liệu (để thấy
được "ghi nửa chừng"). Ghi nhận cho từng biến thể:
  * sập        — ngoại lệ thoát ra ngoài hàm API;
  * sai hợp đồng — kết quả không phải {ok, data, errors} hoặc errors không đúng {code, params} /
                  mã không có trong i18n_errors.MESSAGES;
  * ghi nửa chừng — ok=False nhưng CSDL đổi so với trước khi gọi;
  * lỗi chung  — bị bắt ở lớp ngoài cùng (loi_nap_so_nhap, loi_ngoai_du_kien, ...): đầu vào lạ đã
                 lọt tới một ngoại lệ bất ngờ thay vì bị chặn bằng một lỗi có tên;
  * lộ vết     — errors chứa 'Traceback' hoặc đường dẫn tuyệt đối của máy.
Biến thể được nhận (ok=True) thì chạy tiếp `run_pipeline` và `PRAGMA integrity_check`.
"""

import argparse
import base64
import collections
import copy
import io
import os
import random
import shutil
import sys
import tempfile
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _chung as c  # noqa: E402

import openpyxl  # noqa: E402
import so_nhap  # noqa: E402
from i18n_errors import MESSAGES  # noqa: E402

CAC_LOI_CHUNG = {"loi_nap_so_nhap", "loi_ngoai_du_kien", "error_importing_csv_auto",
                 "error_importing_csv"}

# Giá trị lạ cho ô số (chỉ tiêu, điểm thi, điểm quyết tâm, ...)
SO_LA = ["abc", "-5", "-0.5", "1e308", "1e999", "9" * 40, "٣", "²", "①", "３", "3.7", "3,7", "1.234,5",
         "", "  4  ", "0", "nan", "inf", "-inf", "0x10", "٣٣", "1/2", "10 điểm", "=1+1", "'4", "4.", ".5",
         10 ** 20, -3, 2 ** 53 + 1, 1e-9, True, "TRUE", "N/A", "null", "—"]
TEN_LA = ["", " ", "A" * 5000, "=1+1", "@SUM(A1)", "+cmd", "-2", "Ç̧", "tên\ttab", "tên\nxuống dòng",
          "😀 emoji", "‮giả", "𝔘𝔫𝔦", "x" * 300, "clb' OR '1'='1", "<script>alert(1)</script>",
          "../../etc/passwd", "ａｂｃ", "   khoảng trắng   "]


# ---------------------------------------------------------------- nền
def dung_nen(thu_muc):
    """Dựng CSDL nền (30 em, 6 CLB, 2 buổi) đã nạp + sổ gốc. Trả (sổ, csv_text)."""
    p, clb, hs = c.sinh_bo(30, 6, 2, 777, dqt=True, thu_muc=thu_muc)
    api = c.tao_api(thu_muc, dqt=True)
    r = c.nap_so(api, p)
    assert r["ok"], r
    du, _ = api._doc_so_nhap(c.b64(p))
    return p, so_nhap.thanh_csv(du)


def mo(p):
    return openpyxl.load_workbook(p)


def ghi_b64(wb):
    b = io.BytesIO()
    wb.save(b)
    return base64.b64encode(b.getvalue()).decode()


def hang_du_lieu_hs(ws):
    return range(3, ws.max_row + 1)


def hang_hs_co_du_lieu(ws):
    return [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value]


def cot_theo_tieu_de(ws, ten, hang=1):
    return [i for i in range(1, ws.max_column + 1) if str(ws.cell(hang, i).value or "").strip() == ten]


def cot_phu(ws, ten):
    """Cột con (hàng 2) theo tên: 'CLB' / 'Điểm thi' / 'Điểm quyết tâm'."""
    return cot_theo_tieu_de(ws, ten, 2)


# ---------------------------------------------------------------- toán tử Sổ
def _sheet(wb, rnd):
    return wb["1. CLB"], wb["2. Học sinh"]


def op_clb_o(wb, rnd, cot, gia_tri):
    ws = wb["1. CLB"]
    i = cot_theo_tieu_de(ws, cot)[0]
    ws.cell(rnd.randint(2, ws.max_row), i).value = rnd.choice(gia_tri)


def op_chi_tieu_la(wb, rnd):
    op_clb_o(wb, rnd, "Chỉ tiêu", SO_LA)


def op_suat_uu_tien_la(wb, rnd):
    op_clb_o(wb, rnd, "Suất ưu tiên", SO_LA)


def op_ten_clb_la(wb, rnd):
    op_clb_o(wb, rnd, "Tên CLB", TEN_LA)


def op_ma_clb_la(wb, rnd):
    op_clb_o(wb, rnd, "Mã CLB", TEN_LA + ["clb_tinhoc", "CLB_TINHOC", "clb tinhoc"])


def op_buoi_la(wb, rnd):
    op_clb_o(wb, rnd, "Buổi", ["", "Thứ 9", "thu 2", "THỨ 2", "Thứ  2", "Chủ nhật", "CN", "Thu_2",
                                "Thứ 2 (sáng)", "2", "Monday", "Thứ 2", unicodedata.normalize("NFD", "Thứ 2"),
                                "x" * 200, "Thứ 2/3"])


def op_nfd_ten_clb(wb, rnd):
    ws = wb["1. CLB"]
    for r in range(2, ws.max_row + 1):
        for col in cot_theo_tieu_de(ws, "Tên CLB") + cot_theo_tieu_de(ws, "Mã CLB"):
            v = ws.cell(r, col).value
            if isinstance(v, str) and rnd.random() < 0.7:
                ws.cell(r, col).value = unicodedata.normalize("NFD", v)


def op_nfd_hs(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = [1, 2] + cot_phu(ws, "CLB")
    for r in hang_hs_co_du_lieu(ws):
        for i in cot:
            v = ws.cell(r, i).value
            if isinstance(v, str) and rnd.random() < 0.5:
                ws.cell(r, i).value = unicodedata.normalize("NFD", v)


def op_trung_ten_clb(wb, rnd):
    ws = wb["1. CLB"]
    i = cot_theo_tieu_de(ws, "Tên CLB")[0]
    a, b = rnd.sample(range(2, ws.max_row + 1), 2)
    ws.cell(a, i).value = ws.cell(b, i).value


def op_trung_ma_clb(wb, rnd):
    ws = wb["1. CLB"]
    i = cot_theo_tieu_de(ws, "Mã CLB")[0]
    a, b = rnd.sample(range(2, ws.max_row + 1), 2)
    ws.cell(a, i).value = ws.cell(b, i).value


def op_ma_hs_la(wb, rnd):
    ws = wb["2. Học sinh"]
    rs = hang_hs_co_du_lieu(ws)
    r = rnd.choice(rs)
    ws.cell(r, 1).value = rnd.choice(TEN_LA + [ws.cell(rnd.choice(rs), 1).value, "hs001", "HS 001", 12345, 1.5])


def op_trung_ma_hs(wb, rnd):
    ws = wb["2. Học sinh"]
    a, b = rnd.sample(hang_hs_co_du_lieu(ws), 2)
    ws.cell(a, 1).value = ws.cell(b, 1).value


def op_nv_la(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "CLB")
    r = rnd.choice(hang_hs_co_du_lieu(ws))
    ws.cell(r, rnd.choice(cot)).value = rnd.choice(TEN_LA + ["clb_khong_co", 7, 3.5])


def op_nv_trung(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "CLB")
    r = rnd.choice(hang_hs_co_du_lieu(ws))
    v = ws.cell(r, cot[0]).value
    if v:
        ws.cell(r, cot[1]).value = v


def op_diem_thi_la(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "Điểm thi")
    r = rnd.choice(hang_hs_co_du_lieu(ws))
    ws.cell(r, rnd.choice(cot)).value = rnd.choice(SO_LA)


def op_dqt_la(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "Điểm quyết tâm")
    for _ in range(rnd.randint(1, 6)):
        r = rnd.choice(hang_hs_co_du_lieu(ws))
        ws.cell(r, rnd.choice(cot)).value = rnd.choice(SO_LA)


def op_dqt_rac_ca_cot(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "Điểm quyết tâm")
    rac = rnd.choice(["abc", "٣", "²", "-1", "1e9", "☃"])
    for r in hang_hs_co_du_lieu(ws):
        for i in cot:
            if ws.cell(r, i).value is not None:
                ws.cell(r, i).value = rac


def op_dong_trong(wb, rnd):
    ws = wb["2. Học sinh"]
    for _ in range(rnd.randint(1, 5)):
        ws.insert_rows(rnd.randint(3, ws.max_row))
    ws2 = wb["1. CLB"]
    if rnd.random() < 0.5:
        ws2.insert_rows(rnd.randint(2, ws2.max_row))


def op_xoa_cot_hs(wb, rnd):
    ws = wb["2. Học sinh"]
    ws.delete_cols(rnd.choice([1, 2, 3, 4, 5, 6]), 1)


def op_xoa_cot_clb(wb, rnd):
    ws = wb["1. CLB"]
    ws.delete_cols(rnd.randint(1, 6), 1)


def op_cot_thua(wb, rnd):
    for ws in (wb["1. CLB"], wb["2. Học sinh"]):
        if rnd.random() < 0.6:
            col = ws.max_column + rnd.randint(1, 3)
            ws.cell(1, col).value = rnd.choice(["Ghi chú", "Cột lạ", "Điểm quyết tâm", "Mã CLB", "CLB", "Chỉ tiêu"])
            for r in range(2, min(ws.max_row, 20) + 1):
                ws.cell(r, col).value = rnd.choice(["x", 5, "", None])


def op_o_gop_la(wb, rnd):
    ws = wb["2. Học sinh"]
    r = rnd.choice(hang_hs_co_du_lieu(ws))
    c1 = rnd.randint(1, ws.max_column - 1)
    try:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r + rnd.randint(0, 2),
                       end_column=c1 + rnd.randint(1, 3))
    except Exception:
        pass
    ws2 = wb["1. CLB"]
    if rnd.random() < 0.5:
        r2 = rnd.randint(2, ws2.max_row)
        try:
            ws2.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=rnd.randint(2, 6))
        except Exception:
            pass


def op_khong_em(wb, rnd):
    ws = wb["2. Học sinh"]
    ws.delete_rows(3, ws.max_row)


def op_mot_clb(wb, rnd):
    ws = wb["1. CLB"]
    ws.delete_rows(3, ws.max_row)


def op_khong_clb(wb, rnd):
    ws = wb["1. CLB"]
    ws.delete_rows(2, ws.max_row)


def op_chi_tieu_0(wb, rnd):
    ws = wb["1. CLB"]
    i = cot_theo_tieu_de(ws, "Chỉ tiêu")[0]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value and (rnd.random() < 0.5 or True):
            ws.cell(r, i).value = 0
    j = cot_theo_tieu_de(ws, "Suất ưu tiên")[0]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, j).value = 0


def op_cung_nv1(wb, rnd):
    ws = wb["2. Học sinh"]
    cot = cot_phu(ws, "CLB")
    clb = wb["1. CLB"].cell(2, 1).value
    for r in hang_hs_co_du_lieu(ws):
        for i in cot:
            ws.cell(r, i).value = None
        ws.cell(r, cot[0]).value = clb
    # điểm quyết tâm cũng phải bỏ cho sạch? KHÔNG — để nguyên: đó là một phần của tình huống


def op_bo_sheet(wb, rnd):
    ten = rnd.choice(["1. CLB", "2. Học sinh"])
    del wb[ten]


def op_doi_ten_sheet(wb, rnd):
    wb["1. CLB"].title = rnd.choice(["CLB", "Sheet1", "1. Câu lạc bộ", "1. clb", "1.CLB"])


def op_nhieu_sheet(wb, rnd):
    wb.create_sheet(rnd.choice(["3. Thừa", "2. Học sinh (2)", "Hướng dẫn 2"]))


def op_cong_thuc(wb, rnd):
    ws = wb["2. Học sinh"]
    r = rnd.choice(hang_hs_co_du_lieu(ws))
    ws.cell(r, rnd.choice(cot_phu(ws, "Điểm thi"))).value = rnd.choice(["=1+1", "=A1", "=SUM(1,2)", "=1/0"])
    ws.cell(r, 2).value = "=HYPERLINK(\"http://x\",\"y\")"


def op_rat_lon(wb, rnd):
    ws = wb["2. Học sinh"]
    r0 = ws.max_row
    cot = cot_phu(ws, "CLB")
    clbs = [wb["1. CLB"].cell(i, 1).value for i in range(2, wb["1. CLB"].max_row + 1)]
    for k in range(rnd.randint(500, 3000)):
        ws.cell(r0 + 1 + k, 1).value = "HSX%05d" % k
        ws.cell(r0 + 1 + k, cot[0]).value = rnd.choice(clbs)


OP_SO = {k[3:]: v for k, v in dict(globals()).items() if k.startswith("op_") and callable(v)}
OP_SO.pop("clb_o", None)


# ---------------------------------------------------------------- tệp rác (không phải xlsx hợp lệ)
def tep_rac(rnd, goc_bytes):
    k = rnd.choice(["cat_ngan", "ngau_nhien", "toan_0", "csv_doi_duoi", "zip_rong", "rong", "lap_byte", "html"])
    if k == "cat_ngan":
        return goc_bytes[:rnd.randint(10, len(goc_bytes) - 10)]
    if k == "ngau_nhien":
        return bytes(rnd.getrandbits(8) for _ in range(rnd.randint(1, 5000)))
    if k == "toan_0":
        return b"\x00" * rnd.randint(1, 10000)
    if k == "csv_doi_duoi":
        return "student_id,name\nHS1,a\n".encode()
    if k == "zip_rong":
        return b"PK\x05\x06" + b"\x00" * 18
    if k == "rong":
        return b""
    if k == "lap_byte":
        b = bytearray(goc_bytes)
        for _ in range(rnd.randint(1, 50)):
            b[rnd.randrange(len(b))] = rnd.getrandbits(8)
        return bytes(b)
    return b"<html><body>404</body></html>"


# ---------------------------------------------------------------- toán tử CSV
def mut_csv(rnd, loai, text):
    k = rnd.choice(["bom_giu", "utf16_doc_utf8", "cp1258_doc_utf8", "nfd", "crlf_cr", "cat_giua_dong",
                    "chen_dong_trong", "trung_dong", "xoa_cot_dau", "them_cot", "dau_cham_phay",
                    "so_la", "o_dai", "xoa_tieu_de", "chi_tieu_dau", "ngoac_xuong_dong", "chi_tieu_de",
                    "ma_hs_la"])
    L = text.split("\n")
    if k == "bom_giu":
        return k, "﻿" + text
    if k == "utf16_doc_utf8":
        return k, text.encode("utf-16").decode("utf-8", errors="replace")
    if k == "cp1258_doc_utf8":
        return k, text.encode("cp1258", errors="replace").decode("utf-8", errors="replace")
    if k == "nfd":
        return k, unicodedata.normalize("NFD", text)
    if k == "crlf_cr":
        return k, text.replace("\r\n", "\n").replace("\n", rnd.choice(["\r\n", "\r"]))
    if k == "cat_giua_dong":
        return k, text[:rnd.randint(1, len(text) - 1)]
    if k == "chen_dong_trong":
        for _ in range(rnd.randint(1, 5)):
            L.insert(rnd.randint(1, len(L)), rnd.choice(["", ",,,,", "   "]))
        return k, "\n".join(L)
    if k == "trung_dong":
        i = rnd.randint(1, max(1, len(L) - 2))
        L.insert(i, L[i])
        return k, "\n".join(L)
    if k == "xoa_cot_dau":
        return k, "\n".join(",".join(x.split(",")[1:]) for x in L)
    if k == "them_cot":
        return k, "\n".join(x + ",thua" if x else x for x in L)
    if k == "dau_cham_phay":
        return k, text.replace(",", ";")
    if k == "so_la":
        i = rnd.randint(1, max(1, len(L) - 2))
        parts = L[i].split(",")
        if parts:
            parts[rnd.randrange(len(parts))] = str(rnd.choice(SO_LA)).replace(",", ".").replace("\n", " ")
        L[i] = ",".join(parts)
        return k, "\n".join(L)
    if k == "o_dai":
        i = rnd.randint(1, max(1, len(L) - 2))
        L[i] = L[i] + ",".join(["x" * 100000])
        return k, "\n".join(L)
    if k == "xoa_tieu_de":
        return k, "\n".join(L[1:])
    if k == "chi_tieu_dau":
        return k, text.replace("capacity", "Chi tieu", 1)
    if k == "ngoac_xuong_dong":
        i = rnd.randint(1, max(1, len(L) - 2))
        L[i] = '"' + L[i].replace(",", ',"\n"', 1)
        return k, "\n".join(L)
    if k == "chi_tieu_de":
        return k, text.replace("capacity", "capacity,capacity", 1)
    return k, text.replace("HS0", rnd.choice(["HS0 ", "HS0 ", "hs0", "HS０"]), 3)


# ---------------------------------------------------------------- chạy 1 biến thể
def tao_bien_the(op, seed, nen, thu_muc_nen):
    """Trả (loại, hàm gọi(api) -> kết quả, mô tả)."""
    rnd = random.Random("%s:%s" % (op, seed))
    p_so, csvs = nen
    if op == "tep_rac":
        with open(p_so, "rb") as f:
            goc = f.read()
        data = tep_rac(rnd, goc)
        return "so", (lambda api: api.import_so_nhap(base64.b64encode(data).decode())), "tệp rác"
    if op.startswith("csv:"):
        loai = op.split(":", 1)[1]
        mo_ta, txt = mut_csv(rnd, loai, csvs[loai])
        ham = {"clubs": "import_clubs_csv", "preferences": "import_preferences_csv",
               "test_selection": "import_test_selection_csv"}[loai]
        if rnd.random() < 0.5:
            return "csv", (lambda api: api.import_csv_auto(txt)), "csv_auto/" + mo_ta
        return "csv", (lambda api: getattr(api, ham)(txt)), ham + "/" + mo_ta
    wb = mo(p_so)
    OP_SO[op](wb, rnd)
    b = ghi_b64(wb)
    if rnd.random() < 0.3:
        return "so", (lambda api: api.xem_truoc_so_nhap(b)), "xem_truoc"
    return "so", (lambda api: api.import_so_nhap(b)), "nap"


def kiem_ket_qua(r):
    """Trả danh sách vấn đề về hợp đồng/rò rỉ của một kết quả API."""
    vd = []
    if not isinstance(r, dict) or not {"ok", "data", "errors"} <= set(r):
        return ["sai_hop_dong:khong_phai_ok_data_errors"]
    ds = r["errors"]
    if not isinstance(ds, list):
        return ["sai_hop_dong:errors_khong_la_list"]
    for e in ds:
        if isinstance(e, dict):
            if not isinstance(e.get("code"), str) or not isinstance(e.get("params", {}), dict):
                vd.append("sai_hop_dong:code_params")
            elif e["code"] not in MESSAGES:
                vd.append("sai_hop_dong:ma_la:%s" % e["code"])
            blob = repr(e)
        else:
            blob = str(e)
            vd.append("chuoi_tho")
        if "Traceback" in blob or "/home/" in blob or "\\Users\\" in blob or "/tmp/" in blob:
            vd.append("lo_vet")
    return vd


def chay_mot(op, seed, nen, thu_muc_nen, chi_tiet=False):
    thu_muc = tempfile.mkdtemp(prefix="s2_")
    try:
        shutil.copy(os.path.join(thu_muc_nen, "app.db"), os.path.join(thu_muc, "app.db"))
        api = c.tao_api(thu_muc, dqt=True)
        db = os.path.join(thu_muc, "app.db")
        truoc = c.dau_vet_db(db)
        loai, ham, mo_ta = tao_bien_the(op, seed, nen, thu_muc_nen)
        kq = {"op": op, "seed": seed, "mo_ta": mo_ta, "van_de": []}
        try:
            r = ham(api)
        except BaseException as e:
            kq["van_de"].append("sap:%s" % type(e).__name__)
            kq["chi_tiet"] = repr(e)[:300]
            return kq
        kq["van_de"] += kiem_ket_qua(r)
        ok = isinstance(r, dict) and r.get("ok")
        kq["ok"] = bool(ok)
        sau = c.dau_vet_db(db)
        if not ok:
            if sau != truoc:
                kq["van_de"].append("ghi_nua_chung")
            ma = [e.get("code") for e in r.get("errors", []) if isinstance(e, dict)] if isinstance(r, dict) else []
            kq["ma_loi"] = ma[:3]
            if any(m in CAC_LOI_CHUNG for m in ma):
                kq["van_de"].append("loi_chung:" + ",".join(m for m in ma if m in CAC_LOI_CHUNG))
                kq["chi_tiet"] = repr(r.get("errors"))[:400]
        else:
            # được nhận: chạy pipeline, kiểm toàn vẹn
            try:
                rp = api.run_pipeline(42)
                kq["pipeline_ok"] = bool(rp.get("ok"))
                v2 = kiem_ket_qua(rp)
                kq["van_de"] += ["pipeline_" + x for x in v2]
                if not rp.get("ok"):
                    kq["pipeline_ma"] = [e.get("code") for e in rp.get("errors", []) if isinstance(e, dict)][:3]
            except BaseException as e:
                kq["van_de"].append("sap_pipeline:%s" % type(e).__name__)
                kq["chi_tiet"] = repr(e)[:300]
            if c.toan_ven(db) != "ok":
                kq["van_de"].append("hong_csdl")
        if chi_tiet:
            kq["ket_qua"] = r
        return kq
    finally:
        c.xoa(thu_muc)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--nhanh", action="store_true")
    ap.add_argument("--so-bien-the", type=int, default=0)
    ap.add_argument("--lai", default="")
    a = ap.parse_args()

    thu_muc_nen = tempfile.mkdtemp(prefix="s2_nen_")
    nen = dung_nen(thu_muc_nen)

    if a.lai:
        op, seed = a.lai.rsplit(":", 1)
        kq = chay_mot(op, int(seed), nen, thu_muc_nen, True)
        print(kq)
        return

    ops = sorted(OP_SO) + ["tep_rac", "csv:clubs", "csv:preferences", "csv:test_selection"]
    muc = a.so_bien_the or (len(ops) * 3 if a.nhanh else 640)
    ds = [(ops[i % len(ops)], i // len(ops)) for i in range(muc)]
    tat_ca = []
    for n, (op, seed) in enumerate(ds):
        tat_ca.append(chay_mot(op, seed, nen, thu_muc_nen))
        if (n + 1) % 100 == 0:
            print("  ... %d/%d" % (n + 1, len(ds)), flush=True)

    dem = collections.Counter()
    for k in tat_ca:
        dem["tong"] += 1
        dem["nhan_ok" if k.get("ok") else "tu_choi"] += 1
        if k.get("ok") and k.get("pipeline_ok") is False:
            dem["nhan_nhung_pipeline_tu_choi"] += 1
        for v in set(x.split(":")[0] for x in k["van_de"]):
            dem[v] += 1
    c.in_bang("Nhóm 2 — %d biến thể (Sổ nhập + CSV)" % len(tat_ca),
              ["chỉ số", "số biến thể"], sorted(dem.items()))
    theo_op = collections.defaultdict(lambda: collections.Counter())
    for k in tat_ca:
        o = theo_op[k["op"]]
        o["n"] += 1
        o["ok" if k.get("ok") else "tu_choi"] += 1
        if k["van_de"]:
            o["co_van_de"] += 1
    c.in_bang("Theo toán tử", ["toán tử", "n", "nhận", "từ chối", "có vấn đề"],
              [[op, v["n"], v["ok"], v["tu_choi"], v["co_van_de"]] for op, v in sorted(theo_op.items())])
    bat_thuong = [k for k in tat_ca if k["van_de"]]
    print("\nBiến thể có vấn đề (%d) — tái hiện: python %s --lai OP:SEED" % (len(bat_thuong), os.path.relpath(__file__)))
    nhom = collections.defaultdict(list)
    for k in bat_thuong:
        nhom[(k["op"], tuple(sorted(set(k["van_de"]))))].append(k)
    for (op, vd), ds_k in sorted(nhom.items()):
        print("  %s %s ×%d  vd. %s:%d  %s" % (op, list(vd), len(ds_k), op, ds_k[0]["seed"],
                                              (ds_k[0].get("chi_tiet") or "")[:200]))
    p = c.ghi_json("s2_du_lieu_ban_nhanh" if a.nhanh else "s2_du_lieu_ban",
                   {"may_do": c.may_do(), "tong": dict(dem), "theo_toan_tu": {k: dict(v) for k, v in theo_op.items()},
                    "bien_the_co_van_de": bat_thuong})
    print("\nGhi:", p)
    c.xoa(thu_muc_nen)


if __name__ == "__main__":
    main()
