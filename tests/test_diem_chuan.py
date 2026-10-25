"""Điểm chuẩn từng CLB sau khi chạy (A13.4). Dữ liệu dưới đây là dữ liệu MÔ PHỎNG dựng trong test.

`tinh_diem_chuan` là hàm thuần (nằm ở api_quan_ly.py, không nằm trong thuật toán). Mỗi (CLB, buổi)
cho ra hai tầng x hai mức:
  * tầng 2: điểm quyết tâm HIỆU LỰC thấp nhất trong số em tầng 2 được nhận;
  * tầng 1: điểm thi thấp nhất trong số em tầng 1 được nhận, kèm điểm quỹ 1 của chính em đó;
  * mức = suất thường (`general`) hoặc suất dự trữ (`reserve`) theo `match_results.matched_tier`;
  * còn chỗ ở mức đó thì 0 điểm kèm số chỗ còn; không có ai ở tầng / mức đó thì None.
`get_diem_chuan` chỉ cho quản trị: kiosk không có đường nào gọi nó.
"""

import io

import openpyxl
import pytest

from _wharton import bat_wharton
from api_quan_ly import tinh_diem_chuan
from test_wharton import sinh_bo
from test_wharton_pipeline import _nap_bo_vao_api


def _clb(cap, buoi="", du_tru=0):
    return {"capacity": cap, "reserve_capacity": du_tru, "reserve_group": None, "buoi": buoi}


def _tinh(clubs, thi, don, hl, stb, nhan, **kw):
    """`nhan`: [(sid, club_id, tier)]; buổi lấy theo CLB."""
    rows = [(s, clubs[c]["buoi"], c, t) for s, c, t in nhan]
    return {(r["club_id"], r["buoi"]): r for r in tinh_diem_chuan(
        clubs, thi, don, hl, stb, rows, **kw)}


# ------------------------------------------------------------------ ví dụ tính tay

def test_tang_2_day_ba_em_duoc_nhan_70_50_40_em_bi_tu_choi_30_thi_chuan_la_40():
    clubs = {"X": _clb(3)}
    don = {"X": ["a", "b", "c", "d"]}
    hl = {"X": {"a": 70, "b": 50, "c": 40, "d": 30}}
    stb = {"a": 1, "b": 2, "c": 3, "d": 4}
    r = _tinh(clubs, {"X": {}}, don, hl, stb,
              [("a", "X", "general"), ("b", "X", "general"), ("c", "X", "general")])[("X", "")]
    assert r["suc_chua"] == 3 and r["so_don"] == 4 and r["so_nhan"] == 3 and r["con_cho"] == 0
    assert r["tang2"]["thuong"] == {"diem_quyet_tam": 40, "ma_hoc_sinh_cuoi": "c", "so_nhan": 3}
    assert r["tang1"]["thuong"] is None and r["tang2"]["du_tru"] is None


def test_tang_1_em_cuoi_8_25_diem_thi_va_40_diem_quy_1():
    clubs = {"X": _clb(2)}
    thi = {"X": {"a": 9.0, "b": 8.25, "c": 7.0}}
    hl = {"X": {"a": 100, "b": 40, "c": 90}}
    stb = {"a": 1, "b": 2, "c": 3}
    r = _tinh(clubs, thi, {"X": ["a", "b", "c"]}, hl, stb,
              [("a", "X", "general"), ("b", "X", "general")])[("X", "")]
    assert r["tang1"]["thuong"] == {"diem_thi": 8.25, "diem_quyet_tam": 40,
                                    "ma_hoc_sinh_cuoi": "b", "so_nhan": 2}
    assert r["tang2"]["thuong"] is None


def test_tang_1_bang_diem_thi_thi_em_cuoi_la_em_co_diem_quy_1_thap_hon():
    clubs = {"X": _clb(2)}
    thi = {"X": {"a": 8.0, "b": 8.0}}
    hl = {"X": {"a": 70, "b": 30}}
    r = _tinh(clubs, thi, {"X": ["a", "b"]}, hl, {"a": 2, "b": 1},
              [("a", "X", "general"), ("b", "X", "general")])[("X", "")]
    assert r["tang1"]["thuong"]["ma_hoc_sinh_cuoi"] == "b"
    assert r["tang1"]["thuong"]["diem_quyet_tam"] == 30


def test_clb_con_cho_thi_0_diem_va_ghi_con_bao_nhieu_cho():
    clubs = {"X": _clb(5)}
    hl = {"X": {"a": 60, "b": 20}}
    r = _tinh(clubs, {"X": {}}, {"X": ["a", "b"]}, hl, {"a": 1, "b": 2},
              [("a", "X", "general"), ("b", "X", "general")])[("X", "")]
    assert r["con_cho"] == 3
    assert r["tang2"]["thuong"]["diem_quyet_tam"] == 0
    assert r["tang2"]["thuong"]["con_cho"] == 3


def test_co_suat_du_tru_hai_muc_tach_rieng_theo_matched_tier():
    # sức chứa 3, dự trữ 1: a (dự trữ, 20 điểm), b và c (thường, 90 và 60); d bị từ chối (50).
    clubs = {"X": _clb(3, du_tru=1)}
    hl = {"X": {"a": 20, "b": 90, "c": 60, "d": 50}}
    stb = {"a": 1, "b": 2, "c": 3, "d": 4}
    r = _tinh(clubs, {"X": {}}, {"X": ["a", "b", "c", "d"]}, hl, stb,
              [("a", "X", "reserve"), ("b", "X", "general"), ("c", "X", "general")])[("X", "")]
    assert r["tang2"]["du_tru"]["diem_quyet_tam"] == 20
    assert r["tang2"]["du_tru"]["ma_hoc_sinh_cuoi"] == "a"
    assert r["tang2"]["thuong"]["diem_quyet_tam"] == 60
    assert r["tang2"]["thuong"]["ma_hoc_sinh_cuoi"] == "c"
    assert r["con_cho"] == 0


def test_suat_du_tru_chua_dung_het_thi_muc_du_tru_la_0_diem_con_cho():
    clubs = {"X": _clb(3, du_tru=2)}
    hl = {"X": {"a": 50, "b": 40, "c": 30, "d": 10}}
    r = _tinh(clubs, {"X": {}}, {"X": ["a", "b", "c", "d"]}, hl,
              {"a": 1, "b": 2, "c": 3, "d": 4},
              [("a", "X", "reserve"), ("b", "X", "general"), ("c", "X", "general")])[("X", "")]
    assert r["tang2"]["du_tru"]["diem_quyet_tam"] == 0 and r["tang2"]["du_tru"]["con_cho"] == 1
    # mức thường: 3 chỗ - 1 dự trữ đã dùng = 2 chỗ, nhận 2 -> đầy, chuẩn là 30
    assert r["tang2"]["thuong"]["diem_quyet_tam"] == 30
    assert "con_cho" not in r["tang2"]["thuong"]


def test_nhieu_buoi_tinh_rieng_tung_buoi():
    clubs = {"X": _clb(1, "thu_2"), "Y": _clb(1, "thu_3")}
    hl = {"X": {"a": 80, "b": 20}, "Y": {"a": 10, "b": 90}}
    kq = _tinh(clubs, {"X": {}, "Y": {}}, {"X": ["a", "b"], "Y": ["a", "b"]}, hl,
               {"a": 1, "b": 2}, [("a", "X", "general"), ("b", "Y", "general")])
    assert kq[("X", "thu_2")]["tang2"]["thuong"]["diem_quyet_tam"] == 80
    assert kq[("Y", "thu_3")]["tang2"]["thuong"]["diem_quyet_tam"] == 90
    assert kq[("X", "thu_2")]["buoi"] == "thu_2"


def test_vang_thi_tinh_o_tang_2_voi_diem_hieu_luc():
    # c tick thi mà vắng (không có điểm thi) -> tầng 2 với điểm hiệu lực 0
    clubs = {"X": _clb(2)}
    thi = {"X": {"a": 9.0}}
    hl = {"X": {"a": 100, "c": 0}}
    r = _tinh(clubs, thi, {"X": ["a", "c"]}, hl, {"a": 1, "c": 2},
              [("a", "X", "general"), ("c", "X", "general")])[("X", "")]
    assert r["tang1"]["thuong"]["diem_thi"] == 9.0
    assert r["tang2"]["thuong"]["ma_hoc_sinh_cuoi"] == "c"
    assert r["tang2"]["thuong"]["diem_quyet_tam"] == 0


def test_bang_diem_quyet_tam_thi_em_cuoi_la_em_co_so_boc_tham_lon_hon():
    clubs = {"X": _clb(2)}
    hl = {"X": {"a": 50, "b": 50}}
    r = _tinh(clubs, {"X": {}}, {"X": ["a", "b"]}, hl, {"a": 5, "b": 9},
              [("a", "X", "general"), ("b", "X", "general")])[("X", "")]
    assert r["tang2"]["thuong"]["ma_hoc_sinh_cuoi"] == "b"


def test_clb_khong_ai_duoc_nhan_hai_tang_deu_none_va_con_nguyen_suc_chua():
    clubs = {"X": _clb(2)}
    r = _tinh(clubs, {"X": {}}, {"X": []}, {}, {}, [])[("X", "")]
    assert r["tang1"] == {"thuong": None, "du_tru": None}
    assert r["tang2"] == {"thuong": None, "du_tru": None}
    assert r["con_cho"] == 2 and r["so_nhan"] == 0


def test_hai_lan_tinh_giong_het_nhau_va_khong_sua_dau_vao():
    clubs = {"X": _clb(2)}
    hl = {"X": {"a": 70, "b": 50}}
    vao = (clubs, {"X": {}}, {"X": ["a", "b"]}, hl, {"a": 1, "b": 2},
           [("a", "", "X", "general"), ("b", "", "X", "general")])
    truoc = repr(vao)
    assert tinh_diem_chuan(*vao) == tinh_diem_chuan(*vao)
    assert repr(vao) == truoc


# ------------------------------------------------------------------ bất biến sau khi chạy pipeline thật

def _chay_that(api, seed):
    bo = sinh_bo(seed, n_hs=24, n_clb=5, so_buoi=1 + seed % 2)
    bat_wharton(api)
    _nap_bo_vao_api(api, bo)
    kq = api.run_pipeline(seed=seed)
    assert kq["ok"], kq["errors"]
    return bo


@pytest.mark.parametrize("seed", range(30))
def test_bat_bien_chuan_tang_2_khong_cao_hon_ai_duoc_nhan_va_khong_thap_hon_ai_bi_tu_choi(
        api_factory, seed):
    api = api_factory()
    bo = _chay_that(api, seed)
    kq = api.get_diem_chuan()
    assert kq["ok"], kq["errors"]
    with api._ket_noi_doc() as cur:
        match = [tuple(r) for r in cur.execute(
            "SELECT student_id, buoi, club_id, matched_tier FROM match_results")]
    duoc = {(s, b): c for s, b, c, _ in match}
    hl = bo.diem_hl
    for r in kq["data"]["rows"]:
        cid, b = r["club_id"], r["buoi"]
        tc = r["tang2"]["thuong"]
        if tc is None:
            continue
        # chuẩn <= điểm của mọi em tầng 2 được nhận (ở mức thường)
        thuong2 = [s for s, bb, c, t in match
                   if c == cid and bb == b and t != "reserve" and s not in bo.diem_thi[cid]]
        assert all(tc["diem_quyet_tam"] <= hl[cid].get(s, 0) for s in thuong2)
        if tc.get("con_cho"):
            continue
        # em tầng 2 bị từ chối mà thích CLB này hơn chỗ hiện tại: điểm <= chuẩn
        for s in bo.sids:
            if s in bo.diem_thi[cid] or cid not in bo.prefs[s]:
                continue
            danh_sach = [c for c in bo.prefs[s] if (bo.buoi or {}).get(c, "") == (bo.buoi or {}).get(cid, "")]
            hien = duoc.get((s, b))
            if hien == cid:
                continue
            thich_hon = hien is None or danh_sach.index(cid) < danh_sach.index(hien)
            if thich_hon:
                assert hl[cid].get(s, 0) <= tc["diem_quyet_tam"], (cid, s)


# ------------------------------------------------------------------ API

def test_api_tra_du_lieu_va_ket_qua_cu_khi_dau_van_tay_khac(api):
    bo = _chay_that(api, 3)
    kq = api.get_diem_chuan()
    assert kq["ok"] and kq["data"]["co_ket_qua"] and not kq["data"]["ket_qua_cu"]
    assert {r["club_id"] for r in kq["data"]["rows"]} == set(bo.cids)
    assert all(set(r) >= {"club_id", "buoi", "ten", "suc_chua", "so_don", "so_nhan", "con_cho",
                          "tang1", "tang2"} for r in kq["data"]["rows"])
    # sửa dữ liệu SAU lần chạy (mở khoá rồi đổi một ô điểm) -> kết quả có thể đã cũ
    assert api.mo_khoa_wharton("MO")["ok"]
    sid = bo.sids[0]
    api.admin_save_wharton(sid, {}, {})
    api.create_or_update_club(bo.cids[0], bo.cids[0], bo.cap[bo.cids[0]] + 1, 0, "",
                              (bo.buoi or {}).get(bo.cids[0], ""))
    assert api.get_diem_chuan()["data"]["ket_qua_cu"]


def test_api_tat_tinh_nang_tra_wharton_chua_bat(api):
    kq = api.get_diem_chuan()
    assert not kq["ok"] and kq["errors"][0]["code"] == "wharton_chua_bat"


def test_api_chua_co_ket_qua_tra_rong(api):
    bo = sinh_bo(1, n_hs=6, n_clb=3)
    bat_wharton(api)
    _nap_bo_vao_api(api, bo)
    kq = api.get_diem_chuan()
    assert kq["ok"] and kq["data"]["co_ket_qua"] is False and kq["data"]["rows"] == []


def test_kiosk_khong_co_duong_nao_toi_diem_chuan(api):
    bo = _chay_that(api, 5)
    st = api.get_student_entry_state(bo.sids[0])
    assert st["ok"]
    chu = repr(st["data"]).lower()
    assert "chuan" not in chu and "tang2" not in chu and "tang1" not in chu


# ------------------------------------------------------------------ tệp xuất

def _trang_xuat(api):
    kq = api.export_ket_qua()
    assert kq["ok"], kq["errors"]
    return openpyxl.load_workbook(kq["data"]["path"])


def test_tep_xuat_co_trang_diem_chuan_moi_clb_buoi_mot_dong(api):
    bo = _chay_that(api, 4)
    wb = _trang_xuat(api)
    assert "Điểm chuẩn" in wb.sheetnames
    ws = wb["Điểm chuẩn"]
    noi_dung = [[c.value for c in r] for r in ws.iter_rows()]
    dau = next(i for i, r in enumerate(noi_dung) if r and r[0] == "Mã CLB")
    dong = [r for r in noi_dung[dau + 1:] if r and r[0]]
    assert len(dong) == len(bo.cids)
    cot = noi_dung[dau]
    for ten in ("Tầng 2, suất thường", "Tầng 2, suất dự trữ", "Tầng 1, suất thường",
                "Tầng 1, suất dự trữ"):
        assert ten in cot, cot


def test_tep_xuat_tat_tinh_nang_giong_het_truoc(api):
    # Tắt tính năng: không có trang Điểm chuẩn, các trang còn lại đúng như trước vòng 4.
    api.create_or_update_club("A", "CLB A", 2, 0, "")
    for sid in ("S1", "S2", "S3"):
        api.create_student_if_missing(sid, "Em " + sid)
        api.submit_preferences(sid, ["A"])
    assert api.run_pipeline(seed=7)["ok"]
    wb = _trang_xuat(api)
    assert "Điểm chuẩn" not in wb.sheetnames
    assert len(wb.sheetnames) == 6


def test_tep_xuat_trang_diem_chuan_o_tren_o_cuoi_la_du_lieu_khong_co_cong_thuc(api):
    _chay_that(api, 2)
    ws = _trang_xuat(api)["Điểm chuẩn"]
    for hang in ws.iter_rows():
        for c in hang:
            assert c.data_type != "f"


def test_tep_xuat_khi_ket_qua_cu_co_dong_canh_bao_tren_trang_diem_chuan(api):
    bo = _chay_that(api, 8)
    api.create_or_update_club(bo.cids[0], bo.cids[0], bo.cap[bo.cids[0]] + 1, 0, "",
                              (bo.buoi or {}).get(bo.cids[0], ""))
    ws = _trang_xuat(api)["Điểm chuẩn"]
    chu = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
    assert "có thể đã cũ" in chu


def test_bytes_xuat_khong_la_tep_hong(api):
    _chay_that(api, 6)
    kq = api.export_ket_qua()
    with open(kq["data"]["path"], "rb") as f:
        assert openpyxl.load_workbook(io.BytesIO(f.read())).sheetnames


# ------------------------------------------------------------------ chi tiết từng CLB (dữ liệu cho biểu đồ)

def test_api_chi_tiet_tung_clb_liet_ke_tung_em_tang_2_va_ket_qua(api_factory):
    dem_duoc = 0
    for seed in range(6):
        api = api_factory()
        bo = _chay_that(api, seed)
        for r in api.get_diem_chuan()["data"]["rows"]:
            kq = api.get_diem_chuan_clb(r["club_id"])
            assert kq["ok"], kq["errors"]
            d = kq["data"]
            assert d["club_id"] == r["club_id"] and d["chuan"] == r
            for e in d["em"]:
                assert e["ket_qua"] in ("duoc", "khong")
                assert 0 <= e["diem"] <= 100 and e["student_id"] in bo.sids
                assert e["student_id"] not in bo.diem_thi[r["club_id"]]      # chỉ em tầng 2
                assert e["diem"] == bo.diem_hl[r["club_id"]].get(e["student_id"], 0)
            duoc = {e["student_id"] for e in d["em"] if e["ket_qua"] == "duoc"}
            with api._ket_noi_doc() as cur:
                that = {x[0] for x in cur.execute(
                    "SELECT student_id FROM match_results WHERE club_id = ?", (r["club_id"],))}
            assert duoc == that - set(bo.diem_thi[r["club_id"]])
            dem_duoc += len(duoc)
    assert dem_duoc > 0


def test_api_chi_tiet_em_chi_tinh_em_da_xin_vao_clb_do(api):
    # em đã vào CLB xếp trước CLB này thì không bao giờ bị CLB này từ chối: không tính
    bo = _chay_that(api, 4)
    with api._ket_noi_doc() as cur:
        match = {(r[0], r[1]): r[2] for r in cur.execute(
            "SELECT student_id, buoi, club_id FROM match_results")}
    for cid in bo.cids:
        d = api.get_diem_chuan_clb(cid)["data"]
        buoi = (bo.buoi or {}).get(cid, "")
        for e in d["em"]:
            if e["ket_qua"] == "khong":
                ds = [c for c in bo.prefs[e["student_id"]] if (bo.buoi or {}).get(c, "") == buoi]
                hien = match.get((e["student_id"], buoi))
                assert hien is None or ds.index(cid) < ds.index(hien)
        tong_tang2 = [s for s in bo.sids if cid in bo.prefs[s] and s not in bo.diem_thi[cid]]
        assert len(d["em"]) + d["so_khong_xet"] == len(tong_tang2)


def test_api_chi_tiet_clb_khong_co_tat_tinh_nang_va_ma_sai(api):
    assert api.get_diem_chuan_clb("A")["errors"][0]["code"] == "wharton_chua_bat"
    _chay_that(api, 2)
    kq = api.get_diem_chuan_clb("khong_co_clb_nay")
    assert not kq["ok"]
