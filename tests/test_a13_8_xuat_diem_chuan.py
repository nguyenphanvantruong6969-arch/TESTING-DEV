"""A13.8 lỗi 2.3: `get_diem_chuan` lỗi lúc xuất thì trang "Điểm chuẩn" ghi rõ lý do và kết quả
xuất trả kèm cảnh báo, không xuất trang trống im lặng. Dữ liệu MÔ PHỎNG."""

import openpyxl

import api as api_mod
from _wharton import bat_wharton
from test_wharton import sinh_bo
from test_wharton_pipeline import _nap_bo_vao_api


def _api(tmp_path):
    bo = sinh_bo(77, n_hs=12, n_clb=4)
    api = api_mod.PipelineAPI(str(tmp_path / "x.db"))
    bat_wharton(api)
    _nap_bo_vao_api(api, bo)
    assert api.run_pipeline(seed=3)["ok"]
    return api


def _trang(path):
    wb = openpyxl.load_workbook(path)
    ten = next(n for n in wb.sheetnames if n.startswith("Điểm chuẩn"))
    return "\n".join(str(c.value) for r in wb[ten].iter_rows() for c in r if c.value is not None)


def test_loi_doc_diem_chuan_thi_trang_ghi_ly_do_va_co_canh_bao(tmp_path, monkeypatch):
    api = _api(tmp_path)
    monkeypatch.setattr(api, "get_diem_chuan", lambda *a, **k: {
        "ok": False, "errors": [{"code": "error_reading_results", "params": {}}]})
    kq = api.export_ket_qua(str(tmp_path / "o.xlsx"))
    assert kq["ok"], kq["errors"]
    assert [w["code"] for w in kq["data"]["canh_bao"]] == ["xuat_diem_chuan_loi"]
    chu = _trang(kq["data"]["path"])
    assert "KHÔNG ĐỌC ĐƯỢC ĐIỂM CHUẨN" in chu and "error_reading_results" in chu


def test_khong_loi_thi_khong_canh_bao(tmp_path):
    api = _api(tmp_path)
    kq = api.export_ket_qua(str(tmp_path / "o.xlsx"))
    assert kq["ok"] and kq["data"]["canh_bao"] == []
    assert "KHÔNG ĐỌC ĐƯỢC" not in _trang(kq["data"]["path"])
