"""Vòng sửa lỗi thứ ba của nhánh workspace (hai lượt /code-review cuối) — mỗi test khoá MỘT lỗi.

  A. Quyết định ảnh hưởng kết quả (chạy xếp, trần nhập) đọc cấu hình NGHIÊM: cấu hình đọc không
     được thì dừng với lỗi rõ, không lặng lẽ chạy bằng luật trường học.
  B. Nhập zip: cấu hình trong zip không đọc được thì từ chối, không ghi đè bằng mặc định.
  C. "Sửa lại từ đầu" không xoá điểm theo NGƯỜI (điểm do giáo viên chấm, không gắn với ô tick).
  D. Xoá workspace "mở tạm" không để db_path trỏ vào thư mục đã dời vào _trash.
  E. Chấm điểm theo CLB bị từ chối khi workspace chấm theo người / chỉ theo nhóm; CSV có điểm thì cảnh báo.
  F. Dấu vân tay và run_history.luat_xep dùng luật HIỆU LỰC (danh sách nhóm vô tác dụng không tính).
  G. Mã chết: chon_db_path, khoá bo_diem của tham_so_luat.
  J. Sổ nhập: cột điểm sai chế độ (theo CLB vs theo người) có dữ liệu thì báo lỗi, không lặng lẽ bỏ.
  K. Xuất zip không nén chính tệp zip đang ghi nằm trong thư mục workspace.
  L. update_workspace_config một phần giữ nguyên phần không nhắc tới.
  M. Khởi động nhiều workspace: không mở tạm một workspace hỏng khi còn workspace lành.
  N. Nhóm cứng: nhập CSV không cảnh báo "nhóm không có CLB nào nhận".
  O. Tên trang Excel của tệp xuất dùng chung bộ lọc ký tự với sổ nhập.
"""

import base64
import contextlib
import io
import json
import os
import sqlite3
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

import cau_hinh_ws as ch
import rbda_priority_pipeline as rp
import so_nhap
import workspace as ws
from api import PipelineAPI
from api_xuat import XuatKetQuaMixin


@pytest.fixture
def goc(tmp_path):
    g = tmp_path / "goc"
    g.mkdir()
    return str(g)


def _api(goc, tmp_path, cfg=None, ten="Thử", da_chon=True):
    slug = ws.tao(goc, ten, cfg)
    xuat = tmp_path / ("xuat_" + slug)
    xuat.mkdir(exist_ok=True)
    return PipelineAPI(ws.duong_dan_db(goc, slug), thu_muc_xuat=str(xuat),
                       thu_muc_goc_ws=goc, ws_da_chon=da_chon)


def _cau_hinh_tho(db):
    """Cấu hình đã lưu dạng dict, kết nối ĐÓNG tường minh (Windows không nhả tệp đang mở)."""
    with contextlib.closing(sqlite3.connect(db)) as conn:
        return json.loads(conn.execute("SELECT value FROM workspace_settings").fetchone()[0])


def _hong_cau_hinh(db, gia_tri='{khong phai json'):
    conn = sqlite3.connect(db)
    conn.execute("UPDATE workspace_settings SET value = ? WHERE key = 'config'", (gia_tri,))
    conn.commit()
    conn.close()


def _dung_du_lieu(api):
    assert api.create_or_update_club("a", "A", 1, 0, "")["ok"]
    for sid in ("N1", "N2"):
        assert api.create_student_if_missing(sid, sid)["ok"]
        assert api.submit_preferences(sid, ["a"])["ok"]


def _ma(res):
    return [e["code"] for e in res["errors"]]


# ------------------------------------------------------------------ A. đọc cấu hình nghiêm

def test_run_pipeline_dung_lai_khi_cau_hinh_doc_khong_duoc(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"tier_mode": "strict", "tiers": ["A"]}})
    _dung_du_lieu(api)
    _hong_cau_hinh(api.db_path)
    res = api.run_pipeline(seed=7)
    assert res["ok"] is False
    assert "ws_cau_hinh_hong" in _ma(res)
    conn = sqlite3.connect(api.db_path)
    try:
        assert conn.execute("SELECT COUNT(*) FROM run_history").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM match_results").fetchone()[0] == 0
    finally:
        conn.close()


def test_run_pipeline_khong_ghi_luat_mac_dinh_cho_cau_hinh_hong(goc, tmp_path):
    """Khoá bất biến 4 theo hướng ngược: không bao giờ có dòng lịch sử ghi 'mặc định' sai."""
    api = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}})
    _dung_du_lieu(api)
    cfg = _cau_hinh_tho(api.db_path)
    cfg["rules"]["priority"] = "weighted"          # giá trị của bản phần mềm mới hơn
    _hong_cau_hinh(api.db_path, json.dumps(cfg))
    res = api.run_pipeline(seed=7)
    assert res["ok"] is False and "ws_cau_hinh_hong" in _ma(res)


def test_tran_nhap_nguyen_vong_dung_lai_khi_cau_hinh_hong(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"max_choices": 2}})
    assert api.create_or_update_club("a", "A", 1, 0, "")["ok"]
    assert api.create_student_if_missing("N1", "N1")["ok"]
    _hong_cau_hinh(api.db_path)
    res = api.submit_preferences("N1", ["a"])
    assert res["ok"] is False


def test_khong_co_hang_cau_hinh_van_chay_nhu_mac_dinh(goc, tmp_path):
    api = _api(goc, tmp_path)
    _dung_du_lieu(api)
    conn = sqlite3.connect(api.db_path)
    conn.execute("DELETE FROM workspace_settings")
    conn.commit()
    conn.close()
    assert api.run_pipeline(seed=7)["ok"]


def test_run_pipeline_doc_cau_hinh_mot_lan(goc, tmp_path, monkeypatch):
    api = _api(goc, tmp_path)
    _dung_du_lieu(api)
    dem = {"n": 0}
    goc_doc = ch.doc_cau_hinh

    def dem_doc(*a, **k):
        dem["n"] += 1
        return goc_doc(*a, **k)

    monkeypatch.setattr(ch, "doc_cau_hinh", dem_doc)
    assert api.run_pipeline(seed=7)["ok"]
    assert dem["n"] <= 1, dem["n"]


# ------------------------------------------------------------------ B. nhập zip giữ cấu hình

def test_nhap_zip_tu_choi_khi_cau_hinh_trong_zip_khong_doc_duoc(goc, tmp_path):
    slug = ws.tao(goc, "Nguồn", {"rules": {"tier_mode": "strict", "tiers": ["A"]}})
    db = ws.duong_dan_db(goc, slug)
    cfg = _cau_hinh_tho(db)
    cfg["rules"]["priority"] = "weighted"
    _hong_cau_hinh(db, json.dumps(cfg))
    zp = str(tmp_path / "x.zip")
    ws.xuat_zip(goc, slug, zp)
    truoc = sorted(os.listdir(os.path.join(goc, "workspaces")))
    with pytest.raises(ValueError):
        ws.nhap_zip(goc, zp)
    assert sorted(os.listdir(os.path.join(goc, "workspaces"))) == truoc     # không để lại gì


def test_nhap_zip_cau_hinh_tot_van_nhap_duoc_va_giu_luat(goc, tmp_path):
    slug = ws.tao(goc, "Nguồn", {"rules": {"tier_mode": "strict", "tiers": ["A", "B"]}})
    zp = str(tmp_path / "x.zip")
    ws.xuat_zip(goc, slug, zp)
    moi = ws.nhap_zip(goc, zp)
    assert ch.doc_cau_hinh(ws.duong_dan_db(goc, moi))["rules"]["tiers"] == ["A", "B"]


# ------------------------------------------------------------------ C. sửa lại từ đầu

def test_sua_lai_tu_dau_giu_diem_theo_nguoi(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}})
    _dung_du_lieu(api)
    assert api.submit_person_scores([{"student_id": "N1", "score": 9}])["ok"]
    res = api.reset_student_entry("N1")
    assert res["ok"] and res["data"]["n_scores_removed"] == 0
    diem = {p["student_id"]: p["score"] for p in api.get_person_scoring_list()["data"]["persons"]}
    assert diem["N1"] == 9
    # nguyện vọng thì vẫn bị xoá
    conn = sqlite3.connect(api.db_path)
    assert conn.execute("SELECT COUNT(*) FROM preferences WHERE student_id='N1'").fetchone()[0] == 0
    conn.close()


# ------------------------------------------------------------------ D. xoá workspace mở tạm

def test_xoa_workspace_mo_tam_chuyen_db_path_sang_workspace_con_lai(goc, tmp_path):
    a = ws.tao(goc, "A")
    b = ws.tao(goc, "B")
    os.utime(ws.duong_dan_db(goc, b), (1_000_000_000, 1_000_000_000))      # A là cái gần nhất
    api = PipelineAPI(ws.duong_dan_db(goc, a), thu_muc_goc_ws=goc, ws_da_chon=False)
    assert api.delete_workspace(a, "XOA")["ok"]
    assert api.db_path == ws.duong_dan_db(goc, b)
    assert os.path.isfile(api.db_path)
    assert api.get_dashboard_status()["ok"]


def test_xoa_workspace_khac_khong_dung_vao_db_path(goc, tmp_path):
    a = ws.tao(goc, "A")
    b = ws.tao(goc, "B")
    api = PipelineAPI(ws.duong_dan_db(goc, a), thu_muc_goc_ws=goc, ws_da_chon=False)
    assert api.delete_workspace(b, "XOA")["ok"]
    assert api.db_path == ws.duong_dan_db(goc, a)


# ------------------------------------------------------------------ E. điểm theo CLB sai chế độ

@pytest.mark.parametrize("luat", [{"score_scope": "per_person"},
                                  {"priority": "tier", "tiers": ["A"]}])
def test_cham_diem_theo_clb_bi_tu_choi_khi_khong_dung_diem_theo_clb(goc, tmp_path, luat):
    api = _api(goc, tmp_path, {"rules": luat})
    _dung_du_lieu(api)
    with api._ket_noi_ghi() as cur:
        cur.execute("INSERT INTO club_test_selection VALUES ('N1','a')")
    res = api.submit_club_scores("a", [{"student_id": "N1", "score": 9}])
    assert res["ok"] is False and "scoring_not_per_slot" in _ma(res)
    conn = sqlite3.connect(api.db_path)
    assert conn.execute("SELECT COUNT(*) FROM club_scores").fetchone()[0] == 0
    conn.close()


def test_cham_diem_theo_clb_van_chay_o_che_do_mac_dinh(goc, tmp_path):
    api = _api(goc, tmp_path)
    _dung_du_lieu(api)
    with api._ket_noi_ghi() as cur:
        cur.execute("INSERT INTO club_test_selection VALUES ('N1','a')")
    assert api.submit_club_scores("a", [{"student_id": "N1", "score": 9}])["ok"]


def test_csv_co_diem_canh_bao_khi_workspace_theo_nguoi(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}})
    assert api.create_or_update_club("a", "A", 1, 0, "")["ok"]
    res = api.import_test_selection_csv("student_id,test_club_1,score_1\nN1,a,7\n")
    assert res["ok"], res
    codes = [w["code"] for w in res["data"]["warnings"]]
    assert "csv_score_ignored_mode" in codes
    conn = sqlite3.connect(api.db_path)
    assert conn.execute("SELECT COUNT(*) FROM club_scores").fetchone()[0] == 0
    conn.close()


# ------------------------------------------------------------------ F. luật hiệu lực

def test_luat_hieu_luc_bo_danh_sach_nhom_vo_tac_dung():
    assert rp.luat_xep_de_ghi({"tiers": ["A", "B"]}) is None
    assert rp.luat_xep_de_ghi({"tiers": ["A"], "priority": "score+tier"}) is None
    assert rp.luat_xep_hieu_luc({"tier_mode": "strict", "tiers": ["A"]})["tiers"] == ["A"]
    assert rp.luat_xep_hieu_luc({"priority": "tier", "tiers": ["A"]})["tiers"] == ["A"]
    # chỉ theo nhóm thì điểm không dùng: phạm vi điểm vô tác dụng
    assert rp.luat_xep_hieu_luc({"priority": "tier", "score_scope": "per_person"}
                                )["score_scope"] == "per_slot"


def _dau_van_tay(db, rules):
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    try:
        return rp.dau_van_tay_du_lieu(conn.cursor(), rules)
    finally:
        conn.close()


def test_dau_van_tay_khong_doi_khi_chi_sua_nhom_vo_tac_dung(tmp_path):
    db = str(tmp_path / "a.db")
    rp.init_db(db)
    mac_dinh = _dau_van_tay(db, None)
    assert _dau_van_tay(db, {"tiers": ["A", "B"]}) == mac_dinh
    cung_a = _dau_van_tay(db, {"tier_mode": "strict", "tiers": ["A", "B"]})
    cung_b = _dau_van_tay(db, {"tier_mode": "strict", "tiers": ["B", "A"]})
    assert cung_a != cung_b != mac_dinh


def test_run_full_pipeline_khong_tu_choi_vi_nhom_vo_tac_dung(goc, tmp_path):
    slug = ws.tao(goc, "Rảnh", {"rules": {"tiers": ["A"]}})
    db = ws.duong_dan_db(goc, slug)
    rp.seed_sample_data(db, n_students=20)
    rp.run_full_pipeline(db, seed=1, output_csv_path=str(tmp_path / "x.csv"))


# ------------------------------------------------------------------ G. mã chết

def test_khong_con_ma_chet():
    import main
    assert not hasattr(main, "chon_db_path")
    assert "bo_diem" not in rp.tham_so_luat({"priority": "tier"}, {})


# ------------------------------------------------------------------ J. cột điểm sai chế độ

def _so_voi_diem(api, hang_hs):
    """Sổ mẫu của `api` + một dòng học sinh; trả base64."""
    wb = openpyxl.load_workbook(api.tao_so_nhap_mau()["data"]["path"])
    ten_hs = [t for t in wb.sheetnames if t.startswith("2.")][0]
    wb[ten_hs].append(hang_hs)
    b = io.BytesIO()
    wb.save(b)
    return base64.b64encode(b.getvalue()).decode()


def test_so_theo_clb_nap_vao_workspace_theo_nguoi_bao_loi(goc, tmp_path):
    cu = _api(goc, tmp_path, ten="Cũ")
    assert cu.create_or_update_club("a", "A", 2, 0, "")["ok"]
    # cột mẫu: Mã HS, Họ tên, Nhóm, NV1, Điểm 1, ...
    b64 = _so_voi_diem(cu, ["N1", "Lan", None, "A", 8.5])
    moi = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}}, ten="Mới")
    res = moi.import_so_nhap(b64)
    assert res["ok"] is False
    assert "so_nhap_diem_sai_che_do" in _ma(res)


def test_so_theo_nguoi_nap_vao_workspace_theo_clb_bao_loi(goc, tmp_path):
    nguoi = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}}, ten="Người")
    assert nguoi.create_or_update_club("a", "A", 2, 0, "")["ok"]
    b64 = _so_voi_diem(nguoi, ["N1", "Lan", None, 8.5, "A"])
    cu = _api(goc, tmp_path, ten="Cũ")
    res = cu.import_so_nhap(b64)
    assert res["ok"] is False
    assert "so_nhap_diem_sai_che_do" in _ma(res)


def test_so_cot_diem_sai_che_do_nhung_trong_thi_khong_loi(goc, tmp_path):
    cu = _api(goc, tmp_path, ten="Cũ")
    assert cu.create_or_update_club("a", "A", 2, 0, "")["ok"]
    b64 = _so_voi_diem(cu, ["N1", "Lan", None, "A", None])       # cột điểm có nhưng để trống
    moi = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}}, ten="Mới")
    assert moi.import_so_nhap(b64)["ok"]


# ------------------------------------------------------------------ K. xuất zip trong thư mục workspace

def test_xuat_zip_vao_thu_muc_workspace_khong_nen_chinh_no(goc):
    slug = ws.tao(goc, "Tự nén")
    thu_muc = os.path.dirname(ws.duong_dan_db(goc, slug))
    dich = os.path.join(thu_muc, slug + ".zip")
    ws.xuat_zip(goc, slug, dich)
    with zipfile.ZipFile(dich) as z:
        assert z.testzip() is None
        assert z.namelist() == ["app.db"]
    ws.xuat_zip(goc, slug, dich)                      # xuất lần hai: tệp cũ cũng không được nén vào
    with zipfile.ZipFile(dich) as z:
        assert z.namelist() == ["app.db"]


# ------------------------------------------------------------------ L. cập nhật cấu hình một phần

def test_update_workspace_config_mot_phan_giu_phan_con_lai(goc, tmp_path):
    api = _api(goc, tmp_path, {
        "industry": "company",
        "labels": {"person": {"vi": "Nhân viên", "en": "Employee"}},
        "rules": {"max_choices": 3, "tier_mode": "strict", "tiers": ["A", "B"],
                  "score_scope": "per_person"}})
    res = api.update_workspace_config({"labels": {"slot": {"vi": "Dự án", "en": "Project"}}})
    assert res["ok"], res
    moi = ch.doc_cau_hinh(api.db_path)
    assert moi["labels"]["slot"] == {"vi": "Dự án", "en": "Project"}
    assert moi["labels"]["person"] == {"vi": "Nhân viên", "en": "Employee"}
    assert moi["rules"]["max_choices"] == 3
    assert moi["rules"]["tiers"] == ["A", "B"]
    assert moi["rules"]["score_scope"] == "per_person"
    assert moi["industry"] == "company"


def test_update_workspace_config_chi_doi_ten_khong_dat_lai_luat(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"max_choices": 4}})
    assert api.update_workspace_config({"name": "Tên mới"})["ok"]
    moi = ch.doc_cau_hinh(api.db_path)
    assert moi["name"] == "Tên mới" and moi["rules"]["max_choices"] == 4


def test_update_workspace_config_tu_choi_khi_cau_hinh_cu_hong(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"max_choices": 4}})
    _hong_cau_hinh(api.db_path)
    assert api.update_workspace_config({"name": "X"})["ok"] is False


# ------------------------------------------------------------------ M. khởi động nhiều workspace

@pytest.fixture
def main_mod(tmp_path, monkeypatch):
    import main
    g = tmp_path / "goc_main"
    g.mkdir()
    monkeypatch.setattr(main, "BASE_DIR", str(g))
    monkeypatch.setattr(sys, "argv", ["main.py"])
    return main


def test_khoi_dong_bo_qua_workspace_hong_khi_con_workspace_lanh(main_mod):
    root = main_mod.BASE_DIR
    lanh = ws.tao(root, "Lành")
    hong = ws.tao(root, "Hỏng")
    os.utime(ws.duong_dan_db(root, lanh), (1_000_000_000, 1_000_000_000))
    with open(ws.duong_dan_db(root, hong), "wb") as f:       # hỏng và là cái gần nhất
        f.write(b"rac" * 100)
    db, goc_ws, da_chon = main_mod.khoi_dong_workspace()
    assert db == ws.duong_dan_db(root, lanh)
    assert da_chon is False                                   # còn 2 workspace: phải để người dùng chọn


def test_khoi_dong_moi_workspace_deu_hong_thi_bao_dung_tep_hong(main_mod):
    root = main_mod.BASE_DIR
    hong = ws.tao(root, "Hỏng")
    db = ws.duong_dan_db(root, hong)
    with open(db, "wb") as f:
        f.write(b"rac" * 100)
    with pytest.raises(Exception) as e:
        main_mod.khoi_dong_workspace()
    assert Path(getattr(e.value, "db_path")) == Path(db)


# ------------------------------------------------------------------ N. nhóm cứng không cảnh báo nhóm dự trữ

def test_nhom_cung_khong_canh_bao_nhom_khong_clb_nao_nhan(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"tier_mode": "strict", "tiers": ["Senior", "Junior"]}})
    assert api.create_or_update_club("a", "A", 2, 0, "")["ok"]
    res = api.import_preferences_csv("student_id,reserve_group,pref_1\nN1,Senior,a\nN2,Junior,a\n")
    assert res["ok"], res
    assert "csv_reserve_group_unknown" not in [w["code"] for w in res["data"]["warnings"]]


def test_nhom_mem_van_canh_bao_nhom_khong_clb_nao_nhan(goc, tmp_path):
    api = _api(goc, tmp_path)
    assert api.create_or_update_club("a", "A", 2, 1, "chinh_sach")["ok"]
    res = api.import_preferences_csv("student_id,reserve_group,pref_1\nN1,chinh_sch,a\n")
    assert res["ok"], res
    assert "csv_reserve_group_unknown" in [w["code"] for w in res["data"]["warnings"]]


# ------------------------------------------------------------------ O. tên trang Excel

@pytest.mark.parametrize("ten", ["Dự án/Phòng", "a:b*c?d[e]f", "x" * 50, "  ", ""])
def test_ten_trang_dung_chung_bo_loc_voi_so_nhap(ten):
    mong = so_nhap._ten_sheet(ten) or "Trang"
    assert XuatKetQuaMixin._ten_trang_hop_le(ten, set()) == mong


def test_ten_trang_chi_co_mot_nguon_bo_loc(monkeypatch):
    """Đổi bộ lọc ở so_nhap thì tên trang của tệp xuất đổi theo (không còn bản sao thứ hai)."""
    monkeypatch.setattr(so_nhap, "_ten_sheet", lambda t: "ZZ")
    assert XuatKetQuaMixin._ten_trang_hop_le("bất kỳ", set()) == "ZZ"


# ------------------------------------------------------------------ P. mã lỗi riêng ở mọi chỗ nhập

def _api_hong(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"max_choices": 3}})
    assert api.create_or_update_club("a", "A", 2, 0, "")["ok"]
    assert api.create_student_if_missing("N1", "N1")["ok"]
    _hong_cau_hinh(api.db_path)
    return api


@pytest.mark.parametrize("goi", [
    lambda a: a.submit_preferences("N1", ["a"]),
    lambda a: a.submit_test_selection("N1", ["a"]),
    lambda a: a.submit_club_scores("a", [{"student_id": "N1", "score": 5}]),
    lambda a: a.import_preferences_csv("student_id,pref_1\nN1,a\n"),
    lambda a: a.import_test_selection_csv("student_id,test_club_1,score_1\nN1,a,7\n"),
    lambda a: a.check_data_integrity(),
], ids=["submit_preferences", "submit_test_selection", "submit_club_scores",
        "import_preferences_csv", "import_test_selection_csv", "check_data_integrity"])
def test_cau_hinh_hong_tra_ma_loi_rieng_o_moi_cho_nhap(goc, tmp_path, goi):
    api = _api_hong(goc, tmp_path)
    res = goi(api)
    assert res["ok"] is False
    assert _ma(res) == ["ws_cau_hinh_hong"], res["errors"]


def test_import_so_nhap_cau_hinh_hong_tra_ma_loi_rieng(goc, tmp_path):
    api = _api_hong(goc, tmp_path)
    wb = openpyxl.load_workbook(_api(goc, tmp_path, ten="Mẫu").tao_so_nhap_mau()["data"]["path"])
    b = io.BytesIO()
    wb.save(b)
    res = api.import_so_nhap(base64.b64encode(b.getvalue()).decode())
    assert res["ok"] is False
    assert "ws_cau_hinh_hong" in _ma(res), res["errors"]


def _api_hong_per_person(goc, tmp_path):
    api = _api(goc, tmp_path, {"rules": {"score_scope": "per_person"}}, ten="Người")
    _dung_du_lieu(api)
    _hong_cau_hinh(api.db_path)
    return api


@pytest.mark.parametrize("goi", [
    lambda a: a.tao_so_nhap_mau(),
    lambda a: a.export_du_lieu_dau_vao(),
    lambda a: a.export_hoc_sinh_csv(["N1"]),
    lambda a: a.get_trang_thai_ket_qua(),
    lambda a: a.get_person_scoring_list(),
    lambda a: a.submit_person_scores([{"student_id": "N1", "score": 5}]),
    lambda a: a.xem_truoc_so_nhap(base64.b64encode(b"khong phai xlsx").decode()),
], ids=["tao_so_nhap_mau", "export_du_lieu_dau_vao", "export_hoc_sinh_csv",
        "get_trang_thai_ket_qua", "get_person_scoring_list", "submit_person_scores",
        "xem_truoc_so_nhap"])
def test_cau_hinh_hong_chan_bo_cuc_so_va_tep_xuat(goc, tmp_path, goi):
    """Bố cục sổ nhập / tệp xuất / cách chấm điểm phụ thuộc cấu hình: hỏng thì dừng, không đoán."""
    api = _api_hong_per_person(goc, tmp_path)
    res = goi(api)
    assert res["ok"] is False
    assert "ws_cau_hinh_hong" in _ma(res), res["errors"]
