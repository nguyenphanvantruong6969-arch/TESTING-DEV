"""Lỗi chặn vì hồ sơ điểm quyết tâm hiện gọn, và hai nút xử lý hàng loạt (A13.5;
Playwright + Chromium, qua `browser_host`).

Dữ liệu dưới đây là dữ liệu MÔ PHỎNG: 30 học sinh, 26 em chia thiếu.

  * bấm Chạy xếp: màn hình chỉ có một dòng tóm tắt, 5 dòng em (mỗi em MỘT dòng gộp mọi lỗi của em),
    "và 21 ... khác" và nút "Mở danh sách cần xử lý"; không quá 6 dòng lỗi;
  * API vẫn trả đủ mọi lỗi trong `errors` (hợp đồng {ok, data, errors} không đổi, bất biến 6);
  * nút mở danh sách: sang tab con Điểm quyết tâm, bộ lọc "chỉ hồ sơ có vấn đề" đã bật;
  * "Phóng cho đủ 100 theo tỉ lệ": gõ sai chữ thì không đổi, gõ PHONG thì phóng và liệt kê em bỏ qua;
  * "Tải danh sách cần sửa": tạo tệp Excel đúng số dòng.
"""

import os

import pytest

pytest.importorskip("playwright", reason="chưa cài playwright")
from playwright.sync_api import sync_playwright  # noqa: E402

import browser_host  # noqa: E402
from _wharton import bat_wharton, chen_bids, dong_bids, dung_clb, dung_hs  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

from _trinh_duyet import CHROMIUM  # noqa: E402

N = 30
SAI = ["HS%03d" % i for i in range(1, 27)]          # HS001..HS026 chia thiếu
DUNG = ["HS%03d" % i for i in range(27, 31)]


def _dung_du_lieu(api):
    """A có thi; B, C, D không thi. HS001 toàn số 0 (bỏ qua: tổng 0); HS002 có điểm ở D ngoài
    nguyện vọng (bỏ qua: lỗi khác); HS003 thiếu ở CẢ HAI quỹ; còn lại thiếu 40 ở quỹ 2."""
    bat_wharton(api)
    dung_clb(api, A=(40, ""), B=(40, ""), C=(40, ""), D=(40, ""))
    ds = ["HS%03d" % i for i in range(1, N + 1)]
    dung_hs(api, *ds, tick={s: ["A"] for s in ds})
    assert api.submit_test_selection("HS003", ["A", "B"])["ok"]
    for s in ds:
        assert api.submit_preferences(s, ["A", "B", "C"] if s == "HS002" else ["A", "B", "C", "D"])["ok"]
    assert api.submit_preferences("HS003", ["A", "B", "C", "D"])["ok"]
    dong = []
    for s in ds:
        dong.append((s, "A", 1, 100))
        if s == "HS001":
            dong += [(s, "B", 2, 0), (s, "C", 2, 0), (s, "D", 2, 0)]
        elif s == "HS002":
            dong += [(s, "B", 2, 30), (s, "C", 2, 20), (s, "D", 2, 10)]
        elif s == "HS003":
            dong = dong[:-1] + [(s, "A", 1, 40), (s, "B", 1, 20), (s, "C", 2, 30), (s, "D", 2, 10)]
        elif s in SAI:
            dong += [(s, "B", 2, 30), (s, "C", 2, 20), (s, "D", 2, 10)]
        else:
            dong += [(s, "B", 2, 50), (s, "C", 2, 30), (s, "D", 2, 20)]
    chen_bids(api, dong)


@pytest.fixture
def mo(tmp_path, api):
    if CHROMIUM is None:
        pytest.skip("khong tim thay Chromium")
    dong = []

    def _mo(tab="pipeline"):
        _dung_du_lieu(api)
        api.thu_muc_xuat = str(tmp_path / "TaiXuong")
        os.makedirs(api.thu_muc_xuat, exist_ok=True)
        url = browser_host.serve(api, GOC, "index.html", open_browser=False)
        pw = sync_playwright().start()
        br = pw.chromium.launch(executable_path=CHROMIUM, args=["--no-sandbox"])
        dong.append((br, pw))
        page = br.new_page(locale="vi-VN", viewport={"width": 1280, "height": 1500}, bypass_csp=True)
        page.set_default_timeout(8000)
        loi = []
        page.on("pageerror", lambda e: loi.append(str(e)))
        page.goto(url)
        page.wait_for_selector("#dropZone")
        if tab != "pipeline":
            page.locator('[data-tab="%s"]' % tab).click()
        return page, api, loi

    yield _mo
    for br, pw in dong:
        br.close()
        pw.stop()


def _chay_bi_chan(page):
    page.locator("#btnRun").click()
    page.wait_for_selector("#logPanel:not([hidden]) #logBox .log-dong")


def _sang_tab_wh(page):
    page.locator('[data-tab="admin"]').click()
    page.wait_for_selector("#adminWhartonBlock:not([hidden])", state="attached")
    page.locator('#qlTabs [data-ql-tab="wh"]').click()
    page.wait_for_selector("#whQlBody tr")


# ------------------------------------------------------------------ #1 lỗi chặn gọn

def test_api_van_tra_du_moi_loi_trong_errors(mo):
    page, api, loi = mo()
    kq = api.run_pipeline(seed=1)
    assert not kq["ok"]
    sai = {e["params"]["student_id"] for e in kq["errors"] if e["code"] != "wharton_ho_so_chua_hop_le"}
    assert sai == set(SAI)
    # HS003 thiếu ở cả hai quỹ: đủ hai lỗi, không gộp ở backend
    assert sum(1 for e in kq["errors"] if e["params"].get("student_id") == "HS003") == 2
    assert kq["errors"][0]["code"] == "wharton_ho_so_chua_hop_le" and kq["errors"][0]["params"]["n"] == 26


def test_chay_bi_chan_chi_hien_tom_tat_nam_em_va_con_lai(mo):
    page, api, loi = mo()
    _chay_bi_chan(page)
    dong = page.locator("#logBox .log-dong")
    assert dong.count() == 6                                       # 1 tóm tắt + 5 em
    assert page.locator("#logBox .log-tom-tat").count() == 1
    assert "26" in page.locator("#logBox .log-tom-tat").inner_text()
    em = page.eval_on_selector_all("#logBox .log-em", "ns => ns.map(n => n.textContent.trim())")
    assert [e.split(":")[0] for e in em] == ["HS001", "HS002", "HS003", "HS004", "HS005"]
    # HS003 thiếu ở hai quỹ: MỘT dòng gộp cả hai
    hs003 = next(e for e in em if e.startswith("HS003"))
    assert "quỹ 1 thiếu 40" in hs003 and "quỹ 2 thiếu 60" in hs003
    assert "và 21 học sinh khác" in page.locator("#logBox .log-them").inner_text()
    assert page.locator("#logBox .log-mo-ds").is_visible()
    assert page.locator("#logBox .log-dong").count() <= 6
    assert not loi, loi


def test_tieng_anh_loi_chan_gon_doi_theo(mo):
    page, api, loi = mo()
    _chay_bi_chan(page)
    page.evaluate("window.I18N.setLang('en')")
    page.wait_for_function("document.documentElement.lang === 'en'")
    van_ban = page.locator("#logBox").inner_text()
    assert "26" in page.locator("#logBox .log-tom-tat").inner_text()
    assert "and 21 more" in van_ban and "Open the list to fix" in van_ban
    assert "short by 40" in van_ban
    assert not loi, loi


def test_nut_mo_danh_sach_sang_tab_con_voi_bo_loc_chi_ho_so_co_van_de(mo):
    page, api, loi = mo()
    _chay_bi_chan(page)
    page.locator("#logBox .log-mo-ds").click()
    page.wait_for_selector("#view-admin.is-active")
    page.wait_for_selector('#qlTabs [data-ql-tab="wh"][aria-selected="true"]')
    assert page.locator("#whLocChiLoi").is_checked()
    page.wait_for_selector("#whQlBody tr")
    page.wait_for_function(
        "() => document.getElementById('whQlNhanTrang').textContent.includes('26')")
    assert page.locator("#whQlBody tr:not(.wh-ql-loi)").count() == 0
    assert page.locator("#whQlBody tr").count() == 25             # trang đầu của 26 em sai
    assert not loi, loi


# ------------------------------------------------------------------ #2 hai nút hàng loạt

def test_phong_sai_chu_khong_doi_dung_chu_thi_phong_va_liet_ke_em_bo_qua(mo):
    page, api, loi = mo(tab="admin")
    _sang_tab_wh(page)
    truoc = dong_bids(api)
    page.locator("#btnWhPhong").click()
    assert page.locator("#whPhongHop").is_visible()
    page.locator("#whPhongGo").fill("phong")
    page.locator("#btnWhPhongXn").click()
    page.wait_for_timeout(500)
    assert dong_bids(api) == truoc
    page.locator("#whPhongGo").fill("PHONG")
    page.locator("#btnWhPhongXn").click()
    page.wait_for_selector("#whPhongKetQua:not([hidden])")
    assert dong_bids(api) != truoc
    ds = page.eval_on_selector_all("#whPhongKetQua li", "ns => ns.map(n => n.textContent)")
    assert len(ds) == 2 and ds[0].startswith("HS001") and ds[1].startswith("HS002")
    toast = page.locator("#toastStack").inner_text()
    assert "24 học sinh" in toast and "2 học sinh cần sửa tay" in toast
    canh_bao = api.list_wharton_canh_bao()["data"]
    assert sorted(c["student_id"] for c in canh_bao if c["loai"] == "chua_hop_le") == ["HS001", "HS002"]
    assert not loi, loi


def test_tai_danh_sach_can_sua_tao_tep_dung_so_dong(mo, tmp_path):
    page, api, loi = mo(tab="admin")
    _sang_tab_wh(page)
    truoc = dong_bids(api)
    page.locator("#btnWhTaiCanSua").click()
    thu_muc = api.thu_muc_xuat
    page.wait_for_function(
        "() => document.getElementById('toastStack').textContent.includes('xlsx')")
    tep = [f for f in os.listdir(thu_muc) if f.endswith(".xlsx")]
    assert len(tep) == 1
    from openpyxl import load_workbook
    hang = [r for r in load_workbook(os.path.join(thu_muc, tep[0])).worksheets[0]
            .iter_rows(values_only=True) if r[0] and str(r[0]).startswith("HS")]
    assert len(hang) == 28                                         # HS002 và HS003 mỗi em hai lỗi
    assert {r[0] for r in hang} == set(SAI)
    assert dong_bids(api) == truoc
    assert not loi, loi
