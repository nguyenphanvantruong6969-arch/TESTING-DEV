"""Chụp một Sổ nhập CLB thành dữ liệu so sánh được: ô, độ rộng cột, ô khoá, kiểm dữ liệu,
tô màu có điều kiện, màu ô đầu bảng.

Dùng cho test "mặc định trường học cho ra đúng sổ như trước khi có workspace".
"""

import io

import openpyxl


def chup(nguon) -> dict:
    """`nguon`: đường dẫn, bytes hoặc BytesIO của một tệp .xlsx."""
    if isinstance(nguon, (bytes, bytearray)):
        nguon = io.BytesIO(nguon)
    wb = openpyxl.load_workbook(nguon)
    ra = {"sheets": wb.sheetnames, "active": wb.active.title, "trang": {}}
    for ws in wb.worksheets:
        o = {}
        for hang in ws.iter_rows():
            for c in hang:
                if c.value is not None:
                    o[c.coordinate] = c.value
        mau_dau = {c.coordinate: (c.fill.fgColor.rgb, c.font.bold, c.font.color.rgb if c.font.color else None)
                   for c in ws[1]}
        kiem = sorted(
            (dv.type, str(dv.sqref), dv.formula1, dv.formula2, dv.operator, dv.allow_blank,
             dv.showErrorMessage, dv.errorTitle, dv.error, dv.errorStyle)
            for dv in ws.data_validations.dataValidation)
        dieu_kien = sorted(
            (str(rng.sqref), tuple(r.formula), r.type)
            for rng in ws.conditional_formatting for r in rng.rules)
        ra["trang"][ws.title] = {
            "o": o,
            "rong_cot": {k: round(v.width, 3) for k, v in ws.column_dimensions.items() if v.width},
            "khoa": ws.freeze_panes,
            "cao_hang_1": ws.row_dimensions[1].height,
            "mau_dau": mau_dau,
            "kiem": kiem,
            "dieu_kien": dieu_kien,
            "max": (ws.max_row, ws.max_column),
            "dinh_dang_ma_hs": ws["A2"].number_format,
            # màu nền vài ô nhập mẫu
            "nen": {k: ws[k].fill.fgColor.rgb for k in ("A2", "B2", "C2", "D2", "E2", "F2", "G2")},
        }
    return ra
